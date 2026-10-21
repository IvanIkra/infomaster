"""Teachers' Google Sheets with grades (§5.2, Tablo's §3.1, §3.3): the collector reads, the processing lays out.

The sheets are open to anyone with the link (the setting grades_sheets): each is downloaded as xlsx every 30 minutes,
its title once from the HTML view. In every visible tab the student's own row is found by the name (grades_student);
of a tab the top rows above it (at most TOP_ROWS) and that row are kept (TabView). Where the header ends among the top
rows is the model's to say, so the top rows go to it whole, other students' rows among them (§10); the own name does
not. The layout by Haiku, the 0-10 scale and what counts as a new grade are processing/grades.py's: `on_sheet` gets a
sheet read and says what came of it (Read).

Failures count per sheet: the download, a file that is no workbook (a sheet closed to the link gives a login page),
no own row in any tab (a wrong name, or the row is not there yet: the grades kept are not dropped for it), the layout
(Haiku, the limit). A failed sheet keeps its grades as they were. A sheet failing for 12 hours is told to the System
topic once, and once more when it reads again; the first read of a sheet is told there too.
"""

import asyncio
import datetime as dt
import html
import io
import logging
import re
from collections.abc import Awaitable, Callable
from dataclasses import dataclass, field

import aiohttp

from .. import clock
from ..llm.budget import BudgetExceeded
from ..llm.client import LLMError
from ..settings import GOOGLE_SHEET, Settings

log = logging.getLogger(__name__)

EXPORT_URL = "https://docs.google.com/spreadsheets/d/{id}/export?format=xlsx"
TITLE_URL = "https://docs.google.com/spreadsheets/d/{id}/htmlview"
POLL_EVERY = 30 * 60        # seconds (Tablo's poll_minutes)
TIMEOUT = 120               # seconds
STALE_AFTER = dt.timedelta(hours=12)
TOP_ROWS = 15               # rows of a tab above the own row the model sees: the header is somewhere among them
COURSE_FROM_TITLE = 8       # characters of the sheet's title that name the course when its label is not set
NO_ROW = "своей строки нет ни на одной вкладке: проверьте ФИО на экране «Оценки»"


def sheet_id(url: str) -> str:
    found = GOOGLE_SHEET.search(url)
    if found is None:
        raise ValueError("not a Google Sheets link")
    return found[1]


# --- reading a workbook ------------------------------------------------------------------


def norm(text: str) -> str:
    """Lower case, ё as е, punctuation but "-" as spaces: "ИВАНОВ  Иван," -> "иванов иван"."""
    text = text.lower().replace("ё", "е")
    return " ".join(re.sub(r"[^\w\s-]", " ", text).split())


def matches(cell, student: list[str]) -> bool:
    """The cell starts with the words of the name: "Иванов Иван" finds "ИВАНОВ Иван Петрович"."""
    if not isinstance(cell, str):
        return False
    words = norm(cell).split()
    return len(words) >= len(student) and words[:len(student)] == student


def text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, dt.datetime):
        return value.date().isoformat() if value.time() == dt.time() else value.isoformat(" ", "minutes")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def column(index: int) -> str:
    """0 -> A, 26 -> AA."""
    letters = ""
    index += 1
    while index:
        index, rest = divmod(index - 1, 26)
        letters = chr(65 + rest) + letters
    return letters


@dataclass(frozen=True)
class TabView:
    """What is kept of one tab: its top rows and the student's own row."""

    tab: str
    rows: list[tuple[int, dict[str, str]]]  # the top rows that are not empty: (row number, column letter -> text)
    mine: dict[str, object]       # column letter -> raw value in the student's row, without the name
    name_col: str
    own_row: int                  # its number in the tab

    def header(self, numbers: list[int]) -> list[dict[str, str]]:
        """The rows the model named the header, in tab order."""
        return [cells for number, cells in self.rows if number in numbers]

    def column_header(self, col: str, numbers: list[int]) -> str:
        return " / ".join(row[col] for row in self.header(numbers) if row.get(col))


def read_views(data: bytes, student: str) -> list[TabView]:
    """The visible tabs that have the student's row. ValueError when the data is no workbook."""
    from openpyxl import load_workbook

    target = norm(student).split()
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as e:  # a login page or a broken file: zipfile, openpyxl and XML errors alike
        raise ValueError(f"not a workbook, is the sheet open to anyone with the link? ({type(e).__name__})") from None
    views = []
    try:
        for ws in wb.worksheets:
            if ws.sheet_state != "visible":
                continue
            rows = [list(r) for r in ws.iter_rows(values_only=True)]
            own = next(((i, j) for i, r in enumerate(rows) for j, c in enumerate(r) if matches(c, target)), None)
            if own is None:
                continue
            own_i, name_j = own
            top = [(i + 1, cells) for i, row in enumerate(rows[:min(own_i, TOP_ROWS)])
                   if (cells := {column(j): text(c) for j, c in enumerate(row) if text(c)})]
            mine = {column(j): c for j, c in enumerate(rows[own_i]) if j != name_j and text(c)}
            views.append(TabView(tab=ws.title, rows=top, mine=mine, name_col=column(name_j), own_row=own_i + 1))
    finally:
        wb.close()
    return views


# --- the source ----------------------------------------------------------------------------


@dataclass(frozen=True)
class Sheet:
    """One sheet read: what on_sheet gets."""
    id: str
    url: str
    title: str | None       # as Google names it; None when the title could not be fetched
    course: str             # the short name in front of the labels: the setting's label, else the title's start
    student: str
    views: list[TabView]


@dataclass(frozen=True)
class Read:
    """What on_sheet made of a sheet."""
    first: bool             # the baseline: the first read of the sheet records no new grades
    items: int              # grades in the sheet now
    unread: tuple[str, ...] = ()  # tabs the model could not lay out: skipped until their header changes


@dataclass
class SheetStatus:
    label: str = ""         # the setting's: names the sheet until its title is known
    title: str | None = None
    ok_at: str | None = None
    error: str | None = None
    tabs: int = 0
    items: int = 0
    unread: tuple[str, ...] = ()
    failing_since: dt.datetime | None = field(default=None, repr=False)
    told: bool = False

    @property
    def name(self) -> str:
        return self.title or self.label or "?"

    def public(self) -> dict:
        return {"title": self.title, "ok_at": self.ok_at, "error": self.error, "tabs": self.tabs, "items": self.items,
                "unread": list(self.unread)}


def short_course(label: str, title: str | None, sid: str) -> str:
    return label or (title or sid)[:COURSE_FROM_TITLE].strip()


class GradeSheets:
    def __init__(self, settings: Settings, now: Callable[[], dt.datetime] = clock.now_utc):
        self.settings, self.now = settings, now
        self.on_sheet: Callable[[Sheet], Awaitable[Read]] | None = None
        self.on_poll: Callable[[], Awaitable[None]] | None = None  # after every poll: the new grades go to the forum
        self.notify: Callable[[str], None] | None = None           # a line for the System topic
        self.sheets: dict[str, SheetStatus] = {}                    # by sheet id, the configured sheets only
        self._wake = asyncio.Event()

    def configured(self) -> bool:
        return bool(self.settings.get("grades_student") and self.settings.get("grades_sheets"))

    def status(self) -> dict:
        """The sheets by id; never the links or the name."""
        return {"configured": self.configured(), "sheets": {sid: st.public() for sid, st in self.sheets.items()}}

    def wake(self) -> None:
        """The settings changed: poll now, not in 30 minutes."""
        self._wake.set()

    async def run(self) -> None:
        while True:
            try:
                await self.poll()
            except Exception:
                log.exception("grades poll failed")
            await self._nap(POLL_EVERY)

    async def _nap(self, seconds: float) -> None:
        """Cleared after waking, not before: a change during a poll gives one more poll."""
        try:
            await asyncio.wait_for(self._wake.wait(), seconds)
        except TimeoutError:
            pass
        self._wake.clear()

    async def poll(self) -> None:
        student, sheets = self.settings.get("grades_student"), self.settings.get("grades_sheets")
        ids = [sheet_id(sheet["url"]) for sheet in sheets]
        self.sheets = {sid: self.sheets.get(sid) or SheetStatus() for sid in ids}
        if not student or not sheets:
            return
        async with aiohttp.ClientSession(timeout=aiohttp.ClientTimeout(total=TIMEOUT)) as session:
            for sid, sheet in zip(ids, sheets, strict=True):
                await self._poll_sheet(session, sid, sheet, student)
        if self.on_poll is not None:
            try:
                await self.on_poll()
            except Exception:
                log.exception("posting the new grades failed")

    async def _poll_sheet(self, session: aiohttp.ClientSession, sid: str, sheet: dict, student: str) -> None:
        st = self.sheets[sid]
        st.label = sheet["label"]
        try:
            async with session.get(EXPORT_URL.format(id=sid)) as response:
                response.raise_for_status()
                data = await response.read()
            if st.title is None:  # after the file: a sheet closed to the link has the title of a login page
                st.title = await fetch_title(session, sid)
            views = await asyncio.to_thread(read_views, data, student)
            if not views:
                raise ValueError(NO_ROW)
            read = await self.on_sheet(Sheet(sid, sheet["url"], st.title, short_course(sheet["label"], st.title, sid),
                                             student, views))
        except aiohttp.ClientResponseError as e:  # its text carries the URL, the sheet's id is in it
            self._failed(st, f"HTTP {e.status} {e.message}".strip())
            return
        except (aiohttp.ClientError, TimeoutError, ValueError, OSError) as e:
            self._failed(st, str(e) if str(e) == NO_ROW else f"{type(e).__name__}: {e}")
            return
        except (LLMError, BudgetExceeded) as e:  # the layout; asked again at the next poll
            self._failed(st, str(e))
            return
        except Exception as e:  # one sheet must not stop the others
            log.exception("grades of sheet %s failed", st.name)
            self._failed(st, f"{type(e).__name__}: {e}")
            return
        st.tabs, st.items, st.unread = len(views), read.items, read.unread
        self._recovered(st)
        if read.first:
            self._notify(f"Таблица оценок «{st.name}» подключена: вкладок со своей строкой — {len(views)}, "
                         f"оценок — {read.items}. Новые и изменённые оценки будут приходить в «Оценки».")

    def _failed(self, st: SheetStatus, error: str) -> None:
        st.error = error
        log.warning("grades sheet %s failed: %s", st.name, error)
        now = self.now()
        st.failing_since = st.failing_since or now
        if not st.told and now - st.failing_since >= STALE_AFTER:
            st.told = True
            since = st.failing_since.astimezone(clock.TZ)
            self._notify(f"Таблица оценок «{st.name}» не читается с {since:%d.%m %H:%M}: {error}.\n"
                         "Оценки из неё не обновляются.")

    def _recovered(self, st: SheetStatus) -> None:
        if st.told:
            self._notify(f"Таблица оценок «{st.name}» снова читается.")
        st.ok_at, st.error = clock.local_iso(self.now()), None
        st.failing_since, st.told = None, False

    def _notify(self, line: str) -> None:
        if self.notify is not None:
            self.notify(line)


async def fetch_title(session: aiohttp.ClientSession, sid: str) -> str | None:
    """The sheet's title from its HTML view; None on any failure (asked again at the next poll)."""
    try:
        async with session.get(TITLE_URL.format(id=sid)) as response:
            if response.status != 200:
                return None
            page = await response.text()
    except (aiohttp.ClientError, TimeoutError, UnicodeDecodeError):
        return None
    found = re.search(r"<title>(.*?)</title>", page, re.S)
    if found is None:
        return None
    title = html.unescape(found[1]).strip().removesuffix(" - Google Диск").removesuffix(" - Google Drive").strip()
    return title or None
