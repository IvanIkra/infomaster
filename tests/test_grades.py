"""Grades on synthetic sheets shaped like the real ones (names are made up); Tablo's tests/test_grades.py and ours."""

import datetime as dt
import io

import pytest
from aiohttp import web
from openpyxl import Workbook
from test_http import TWIN, env, owner

from uni_master import clock
from uni_master.collectors import grades as collector
from uni_master.collectors.grades import GradeSheets, Read, Sheet
from uni_master.llm.budget import BudgetExceeded
from uni_master.llm.client import BadAnswer
from uni_master.outputs import api, miniapp
from uni_master.outputs.grades import GradeTopic, line, mark
from uni_master.processing import grades as g
from uni_master.processing.grades import Grades, Item
from uni_master.settings import Settings, SettingsError

URL = "https://docs.google.com/spreadsheets/d/abc123/edit#gid=0"
OTHER = "https://docs.google.com/spreadsheets/d/zzz999/edit"


def workbook(hw_score=5.9, quiz=8, hidden_tab=True, inserted=False, me="Иванов Иван Сергеевич", classmate=10,
             total="Балл") -> bytes:
    wb = Workbook()
    # Like the discrete maths sheet: max in row 1, per-problem points, a total and a bonus column.
    ws = wb.active
    ws.title = "ДЗ-1А"
    ws.append([None, "Max за обыкновенные / сложные:", None, 7, 0])
    ws.append(["Группа", "Студент", "Группа*", "1а", "1б", total, "Балл*"])
    ws.append([260, "Образцов Нагляд Примерович", None, 1, 0, 1, 0])
    ws.append([None, "Петров Пётр Петрович", None, 1, 1, 7, 0])
    ws.append([None, me, None, 1, 0.7, hw_score, 1])
    # Like the calculus sheet: two-row header with a grade after each quiz; `inserted` puts a column before them.
    ws = wb.create_sheet("Квизы")
    extra = [None] if inserted else []
    ws.append(["Первый семестр"])
    ws.append([None, *extra, "Средний балл", "Квиз 1", None, "Оценка", "Квиз 2", "Оценка"])
    ws.append(["ФИО", *extra, None, "№1", "№2", None, None, None])
    ws.append(["сидорова анна", *extra, 9, 1, 1, classmate, None, 0])  # lower case: Tablo took it for a header
    ws.append([me, *extra, 7, 0.6, 1, quiz, None, 0])
    if hidden_tab:
        ws = wb.create_sheet("old")
        ws.append(["ФИО", "ДЗ"])
        ws.append(["Иванов Иван", 3])
        ws.sheet_state = "hidden"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def spec(prompt: str) -> dict:
    """The layout Haiku would give, by the tab and where its columns are in the header."""
    tab = prompt.split('"')[1]
    if tab == "ДЗ-1А":
        return {"skip": False, "reason": "", "header_rows": [1, 2], "items": [
            {"column": "F", "label": "HW 1A", "max": 7, "bonus_column": "G"}]}
    shift = 1 if "C=Средний балл" in prompt else 0
    quiz1, quiz2 = "EF"[shift], "GH"[shift]
    return {"skip": False, "reason": "", "header_rows": [1, 2, 3, 99], "items": [  # 99: not a row it was shown
        {"column": quiz1, "label": "Quiz 1", "max": 10, "bonus_column": None},
        {"column": quiz2, "label": "Quiz 2", "max": 10, "bonus_column": None}]}


class FakeLLM:
    available = True

    def __init__(self):
        self.prompts, self.error = [], None

    async def json(self, task, system, prompt, schema, *, max_tokens):
        assert task == "grades" and "{budget}" not in system
        self.prompts.append(prompt)
        if self.error is not None:
            raise self.error
        return spec(prompt)


@pytest.fixture
def settings(store):
    settings = Settings(store)
    settings.set("grades_student", "Иванов Иван")
    settings.set("grades_sheets", [{"url": URL, "label": "DM"}])
    return settings


@pytest.fixture
def llm():
    return FakeLLM()


@pytest.fixture
def grades(store, settings, llm, now):
    return Grades(store, settings, llm, now=now)


def views(**kwargs):
    return collector.read_views(workbook(**kwargs), "Иванов Иван")


def sheet(**kwargs) -> Sheet:
    return Sheet("abc123", URL, "ДМ ПИ", "DM", "Иванов Иван", views(**kwargs))


# --- reading ---------------------------------------------------------------------------------

def test_read_views_keeps_the_top_rows_and_the_own_row():
    found = {v.tab: v for v in views()}
    assert set(found) == {"ДЗ-1А", "Квизы"}  # the hidden tab is skipped
    dz = found["ДЗ-1А"]
    assert [row for row, _ in dz.rows] == [1, 2, 3, 4] and dz.own_row == 5 and dz.mine["F"] == 5.9
    assert dz.header([2]) == [{"A": "Группа", "B": "Студент", "C": "Группа*", "D": "1а", "E": "1б", "F": "Балл",
                               "G": "Балл*"}]
    prompt = g.layout_prompt(dz)
    assert "row 3: A=260 | B=Образцов Нагляд Примерович" in prompt  # the model tells the header from students
    assert prompt.endswith("student row (row 5): D=1 | E=0.7 | F=5.9 | G=1")
    for view in found.values():
        assert "Иванов" not in g.layout_prompt(view)  # the own name is not sent


def test_name_matching_handles_case_and_yo():
    assert collector.matches("ИВАНОВ  Иван Сергеевич", ["иванов", "иван"])
    assert collector.matches("Иванов Иван", ["иванов", "иван"])
    assert not collector.matches("Иванова Ирина", ["иванов", "иван"])
    assert collector.norm("Пётр") == "петр"
    assert collector.column(0) == "A" and collector.column(26) == "AA"


def test_not_a_workbook():
    with pytest.raises(ValueError, match="not a workbook"):
        collector.read_views(b"<html>Sign in</html>", "Иванов Иван")


def test_extract_normalises_to_ten_and_keeps_bonus():
    view = next(v for v in views() if v.tab == "ДЗ-1А")
    [item] = g.extract("sid", "DM", view, spec('Tab: "ДЗ-1А"'))  # header rows 1 and 2
    assert item.label == "DM HW 1A" and item.value == pytest.approx(8.43) and item.bonus == 1
    assert (item.raw, item.max) == (5.9, 7) and item.key == "sid|ДЗ-1А|F|Балл"
    assert g.budget("DM") == 20 and g.budget("Very long course") == 8


def test_prompt_spells_out_cyrillic_letter_order():
    assert "в→C" in g.LAYOUT_SYSTEM and '"HW 1C"' in g.LAYOUT_SYSTEM


# --- the layout --------------------------------------------------------------------------------

def quizzes(**kwargs):
    return next(view for view in views(**kwargs) if view.tab == "Квизы")


async def test_layout_stands_while_the_header_rows_it_named_stay(grades, llm):
    assert (await grades.layout("abc123", quizzes(), 20))["header_rows"] == [1, 2, 3]
    await grades.layout("abc123", quizzes(quiz=6), 20)          # the own grade
    await grades.layout("abc123", quizzes(classmate=4), 20)     # another student's, in the top rows
    assert len(llm.prompts) == 1
    await grades.layout("abc123", quizzes(), 12)                # another budget: another label
    await grades.layout("zzz999", quizzes(), 20)                # another sheet
    assert len(llm.prompts) == 3
    dz = views()[0]
    await grades.layout("abc123", dz, 20)
    await grades.layout("abc123", views(total="Итог ДЗ")[0], 20)  # the header changed
    assert len(llm.prompts) == 5


async def test_an_answer_the_model_cannot_give_skips_the_tab_until_it_changes(grades, llm):
    llm.error = BadAnswer("answer was cut at max_tokens")
    assert (await grades.layout("abc123", quizzes(), 20))["skip"]
    llm.error = None
    assert (await grades.layout("abc123", quizzes(), 20))["skip"] and len(llm.prompts) == 1
    assert not (await grades.layout("abc123", quizzes(classmate=4), 20))["skip"]  # no header known: any change asks
    assert len(llm.prompts) == 2


async def test_a_failed_layout_records_nothing(grades, llm):
    llm.error = BudgetExceeded("the daily LLM limit is spent")
    with pytest.raises(BudgetExceeded):
        await grades.apply(sheet())
    assert grades.db.execute("select count(*) n from grade_sheets").fetchone()["n"] == 0
    llm.error = None
    assert await grades.apply(sheet()) == Read(True, 2)  # Quiz 2 is a placeholder zero


# --- new grades --------------------------------------------------------------------------------

def item(key, label, value, bonus=None, position=0, tab="T", header=None):
    return Item(key, tab, header or key, label, value, value, 10, bonus, position)


def test_first_read_is_a_baseline_then_changes_are_new(grades, now):
    assert grades.record("abc123", "ДМ", [item("k1", "DM HW 1A", 8.43), item("k2", "DM Quiz 1", 0)])
    assert grades.new() == []
    assert [i["label"] for i in grades.items("abc123")] == ["DM HW 1A"]  # the zero is a placeholder

    now.moment += dt.timedelta(days=1)
    assert not grades.record("abc123", "ДМ", [item("k1", "DM HW 1A", 9.0), item("k2", "DM Quiz 1", 8, position=1),
                                              item("k3", "DM Quiz 2", 0, position=2)])
    new = grades.new()
    assert {(x["label"], x["value"]) for x in new} == {("DM HW 1A", 9.0), ("DM Quiz 1", 8)}
    assert set(new[0]) == {"key", "label", "value", "bonus", "seen_at"}
    events = {e["label"]: e for e in grades.waiting()}
    assert events["DM HW 1A"]["previous"] == 8.43 and events["DM Quiz 1"]["previous"] is None
    assert events["DM HW 1A"]["course"] == "DM" and events["DM HW 1A"]["url"] == URL

    now.moment += dt.timedelta(days=8)
    assert grades.new() == [] and grades.waiting() == []  # out of the window


def test_the_newest_event_of_a_work_is_given(grades, now):
    grades.record("abc123", None, [])
    for value in (5, 6, 7):
        now.moment += dt.timedelta(hours=1)
        grades.record("abc123", None, [item("k1", "DM HW 1", value)])
    assert [(x["label"], x["value"]) for x in grades.new()] == [("DM HW 1", 7)]


async def test_a_moved_column_is_no_new_grade(grades, now):
    """A column inserted before the quizzes shifts their letters: Tablo would post them again."""
    await grades.apply(sheet())
    now.moment += dt.timedelta(hours=1)
    await grades.apply(sheet(inserted=True))
    assert grades.new() == []
    assert [i["label"] for i in grades.items("abc123")] == ["DM HW 1A", "DM Quiz 1"]
    await grades.apply(sheet(inserted=True, quiz=9))
    assert [(x["label"], x["value"]) for x in grades.new()] == [("DM Quiz 1", 9)]


async def test_a_grade_that_leaves_and_comes_back_is_no_new_grade(grades, now):
    """A tab hidden for a while, a row gone from a tab: Tablo deleted the grades and posted them again on return."""
    await grades.apply(sheet())
    now.moment += dt.timedelta(hours=1)
    whole = sheet()
    hidden = Sheet(whole.id, whole.url, whole.title, whole.course, whole.student,
                   [view for view in whole.views if view.tab != "Квизы"])
    assert await grades.apply(hidden) == Read(False, 1)
    assert [i["label"] for i in grades.items("abc123")] == ["DM HW 1A"]
    await grades.apply(sheet())
    assert grades.new() == [] and len(grades.items("abc123")) == 2
    await grades.apply(hidden)
    await grades.apply(sheet(quiz=9))  # back with another grade: that one is new, with the grade before it
    assert [(e["label"], e["value"], e["previous"]) for e in grades.waiting()] == [("DM Quiz 1", 9, 8)]


def test_a_work_listed_twice_counts_once(grades):
    view = views()[0]
    twice = {"skip": False, "reason": "", "items": [
        {"column": "F", "label": "HW 1A", "max": 7, "bonus_column": None},
        {"column": "F", "label": "HW 1A again", "max": 7, "bonus_column": None}]}
    items = g.extract("abc123", "DM", view, twice)
    assert [i.label for i in items] == ["DM HW 1A"]


async def test_a_tab_the_model_cannot_lay_out_is_told(grades, llm):
    llm.error = BadAnswer("answer was cut at max_tokens")
    read = await grades.apply(sheet())
    assert read == Read(True, 0, ("ДЗ-1А", "Квизы"))


async def test_a_wide_tab_gets_room_for_its_answer(grades, llm):
    asked = []

    async def json(task, system, prompt, schema, *, max_tokens):
        asked.append(max_tokens)
        return spec(prompt)

    llm.json = json
    view = views()[0]
    wide = collector.TabView(view.tab, [(1, {collector.column(i): f"Sep {i}" for i in range(2, 202)})], view.mine,
                             "B", 2)
    await grades.layout("abc123", view, 20)
    await grades.layout("zzz999", wide, 20)
    assert asked == [g.MAX_TOKENS, 200 * g.TOKENS_PER_WORK]  # the own row's columns are among the 200


# --- the owner's settings ----------------------------------------------------------------------

def test_settings_check_the_sheets():
    check = g.SPECS["grades_sheets"][1]
    assert check([{"url": f" {URL} ", "label": " DM "}]) == [{"url": URL, "label": "DM"}]
    for wrong in ([{"url": "https://example.com/spreadsheets/d/abc"}], [{"url": URL}, {"url": URL + "x&gid=1"}],
                  [{"url": URL, "label": "Discrete Mathematics"}], [{"url": URL, "extra": 1}], {"url": URL}):
        with pytest.raises(SettingsError):
            check(wrong)
    assert g.SPECS["grades_student"][1]("  Иванов   Иван ") == "Иванов Иван"


def test_configure_forgets_removed_sheets_and_everything_for_a_new_name(grades, settings):
    woken = []
    grades.wake = lambda: woken.append(1)
    grades.record("abc123", None, [item("k1", "DM HW", 8)])
    grades.configure({"grades_sheets": [{"url": URL, "label": "DM"}, {"url": OTHER, "label": "Calc"}]})
    grades.record("zzz999", None, [item("k2", "Calc Quiz", 7)])
    assert grades.items("abc123") and grades.items("zzz999") and woken == [1]

    grades.configure({"grades_sheets": [{"url": OTHER, "label": "Calc"}]})
    assert not grades.items("abc123") and grades.items("zzz999")
    grades.configure({"grades_student": "Иванов  Иван"})  # the same name: nothing forgotten
    assert grades.items("zzz999")
    grades.configure({"grades_student": "Петров Пётр"})
    assert not grades.items("zzz999")
    assert grades.db.execute("select count(*) n from grade_sheets").fetchone()["n"] == 0

    with pytest.raises(SettingsError):
        grades.configure({"grades_window_days": 0, "grades_student": "Кто-то"})
    assert settings.get("grades_student") == "Петров Пётр"  # nothing saved when one value is wrong
    with pytest.raises(SettingsError):
        grades.configure({"llm_daily_limit_usd": 1})


async def test_a_sheet_dropped_while_read_records_nothing(grades, settings):
    read = sheet()
    settings.set("grades_sheets", [])
    assert await grades.apply(read) == Read(False, 0)
    settings.set("grades_sheets", [{"url": URL, "label": "DM"}])
    settings.set("grades_student", "Петров Пётр")
    assert await grades.apply(read) == Read(False, 0)
    assert grades.db.execute("select count(*) n from grade_items").fetchone()["n"] == 0


# --- the topic ------------------------------------------------------------------------------------

def test_lines():
    assert mark(8.43) == "8.4" and mark(9.0) == "9" and mark(9.96) == "10"
    event = {"label": "DM HW 1A", "value": 8.43, "raw": 5.9, "max": 7, "bonus": 1.0, "previous": None}
    assert line(event, "ru") == "DM HW 1A — <b>8.4</b> +1 (5.9 из 7)"
    assert line(event, "en") == "DM HW 1A — <b>8.4</b> +1 (5.9 of 7)"
    event = {"label": "Calc <Quiz> 2", "value": 6, "raw": 6, "max": 10, "bonus": None, "previous": 5.5}
    assert line(event, "ru") == "Calc &lt;Quiz&gt; 2 — <b>6</b> (было 5.5)"
    assert line(event, "en") == "Calc &lt;Quiz&gt; 2 — <b>6</b> (was 5.5)"


async def test_topic_posts_a_message_per_sheet_and_keeps_failures(grades, forum, now):
    topic = GradeTopic(forum, grades.store, grades, now=now)
    grades.configure({"grades_sheets": [{"url": URL, "label": "DM"}, {"url": OTHER, "label": ""}]})
    grades.record("abc123", "ДМ ПИ", [])
    grades.record("zzz999", "Матанализ ПИ", [])
    grades.record("abc123", "ДМ ПИ", [item("k1", "DM HW 1", 8)])
    grades.record("zzz999", "Матанализ ПИ", [item("k2", "Матанали Quiz 1", 7)])

    async def down(*args, **kwargs):
        raise ConnectionError("Telegram is down")

    post, forum.post = forum.post, down
    await topic.publish()
    assert len(grades.waiting()) == 2
    forum.post = post
    await topic.publish()
    assert forum.texts("grades") == [
        f'📊 <b>Оценки · <a href="{URL}">ДМ ПИ</a></b> #DM\nDM HW 1 — <b>8</b>',
        '📊 <b>Оценки · <a href="https://docs.google.com/spreadsheets/d/zzz999/edit">Матанализ ПИ</a></b> #Матанали\n'
        "Матанали Quiz 1 — <b>7</b>"]
    assert grades.waiting() == []
    await topic.publish()
    assert len(forum.texts("grades")) == 2


async def test_a_failure_halfway_posts_nothing_twice(grades, forum, now, monkeypatch):
    from uni_master.outputs import grades as topic_module

    monkeypatch.setattr(topic_module, "MESSAGE_LIMIT", 140)  # the header and two lines
    topic = GradeTopic(forum, grades.store, grades, now=now)
    grades.record("abc123", "ДМ", [])
    grades.record("abc123", "ДМ", [item(f"k{i}", f"DM HW {i}", 8, position=i) for i in range(6)])
    post, calls = forum.post, []

    async def flaky(*args, **kwargs):
        calls.append(1)
        if len(calls) == 2:
            raise ConnectionError("Telegram is down")
        return await post(*args, **kwargs)

    forum.post = flaky
    await topic.publish()
    assert len(forum.texts("grades")) == 1 and len(grades.waiting()) == 4
    forum.post = post
    await topic.publish()
    lines = [text for message in forum.texts("grades") for text in message.split("\n")[1:]]
    assert lines == [f"DM HW {i} — <b>8</b>" for i in range(6)] and len(forum.texts("grades")) == 3


# --- the collector ----------------------------------------------------------------------------------

@pytest.fixture
async def google(aiohttp_server, monkeypatch):
    """Google Sheets: the export of a sheet and its HTML view; `state["data"]` is the workbook, or an error code."""
    state = {"data": workbook(), "titles": 0}

    async def export(request):
        if isinstance(state["data"], int):
            return web.Response(status=state["data"])
        return web.Response(body=state["data"])

    async def title(request):
        state["titles"] += 1
        return web.Response(text="<html><title>ДМ &quot;ПИ&quot; - Google Диск</title></html>",
                            content_type="text/html")

    app = web.Application()
    app.router.add_get("/d/{id}/export", export)
    app.router.add_get("/d/{id}/htmlview", title)
    server = await aiohttp_server(app)
    base = f"http://{server.host}:{server.port}"
    monkeypatch.setattr(collector, "EXPORT_URL", base + "/d/{id}/export")
    monkeypatch.setattr(collector, "TITLE_URL", base + "/d/{id}/htmlview")
    return state


async def test_poll_end_to_end(google, grades, settings, llm, forum, now):
    source = GradeSheets(settings, now=now)
    topic = GradeTopic(forum, grades.store, grades, now=now)
    source.on_sheet, source.on_poll = grades.apply, topic.publish
    system = []
    source.notify = lambda line: system.append(line.ru)  # a Phrase: the forum picks its language

    await source.poll()  # the baseline
    status = source.status()["sheets"]["abc123"]
    assert status["title"] == 'ДМ "ПИ"' and (status["tabs"], status["items"], status["error"]) == (2, 2, None)
    assert system == ['Таблица оценок «ДМ "ПИ"» подключена: вкладок со своей строкой — 2, оценок — 2. '
                      "Новые и изменённые оценки будут приходить в «Оценки»."]
    assert forum.texts("grades") == []

    google["data"] = workbook(hw_score=6.3, quiz=9)
    now.moment += dt.timedelta(hours=1)
    await source.poll()
    assert forum.texts("grades") == [f'📊 <b>Оценки · <a href="{URL}">ДМ &quot;ПИ&quot;</a></b> #DM\n'
                                     "DM HW 1A — <b>9</b> +1 (6.3 из 7; было 8.4)\nDM Quiz 1 — <b>9</b> (было 8)"]
    assert len(llm.prompts) == 2 and google["titles"] == 1  # one layout per tab, the title once
    assert len(system) == 1


async def test_poll_failures(google, grades, settings, now):
    source = GradeSheets(settings, now=now)
    source.on_sheet = grades.apply
    system = []
    source.notify = lambda line: system.append(line.ru)  # a Phrase: the forum picks its language
    await source.poll()
    assert grades.items("abc123")

    google["data"] = 404
    await source.poll()
    assert source.status()["sheets"]["abc123"]["error"] == "HTTP 404 Not Found"
    google["data"] = b"<html>sign in</html>"
    await source.poll()
    assert "not a workbook" in source.status()["sheets"]["abc123"]["error"]
    google["data"] = workbook(me="Петров Иван Ильич")  # the row is gone: the grades stay
    await source.poll()
    assert source.status()["sheets"]["abc123"]["error"] == collector.NO_ROW and grades.items("abc123")
    assert len(system) == 1  # only "подключена"

    now.moment += dt.timedelta(hours=13)
    await source.poll()
    await source.poll()
    assert system[1].startswith("Таблица оценок «ДМ \"ПИ\"» не читается с 20.10 12:00: своей строки нет") \
        and len(system) == 2
    google["data"] = workbook()
    await source.poll()
    assert system[2:] == ['Таблица оценок «ДМ "ПИ"» снова читается.'] and grades.new() == []
    assert source.status()["sheets"]["abc123"]["error"] is None


async def test_poll_reports_errors_without_raising(settings, grades, monkeypatch, now):
    monkeypatch.setattr(collector, "EXPORT_URL", "http://127.0.0.1:9/{id}")
    source = GradeSheets(settings, now=now)
    source.on_sheet = grades.apply
    await source.poll()
    assert source.status()["sheets"]["abc123"]["error"]
    settings.set("grades_sheets", [])
    await source.poll()
    assert source.status() == {"configured": False, "sheets": {}}


# --- the ports ----------------------------------------------------------------------------------------

async def test_api_and_miniapp(aiohttp_client, grades, settings, budget, journal, store, schedule, merged, now):

    grades.record("abc123", "ДМ ПИ", [])
    now.moment += dt.timedelta(hours=1)
    grades.record("abc123", "ДМ ПИ", [item("k1", "DM HW 1", 8, 1.0)])
    status = {"configured": True, "sheets": {"abc123": {"title": "ДМ ПИ", "ok_at": "2026-10-20T13:00:00",
                                                        "error": None, "tabs": 1, "items": 1}}, "llm": True, "new": 1}
    sources = {**TWIN, "grades": lambda: status}

    client = await aiohttp_client(api.create_app(env(), budget, sources, None, merged, grades))
    body = await (await client.get("/api/grades", headers={"Authorization": "Bearer tok"})).json()
    assert body == {"items": [{"key": "k1", "label": "DM HW 1", "value": 8, "bonus": 1,
                               "seen_at": clock.local_iso(now())}], "window_days": 7}
    assert (await client.get("/api/grades")).status == 401

    client = await aiohttp_client(miniapp.create_app(env(), budget, journal, schedule.inbox, None, sources,
                                                     schedule.subjects, schedule, None, merged, settings,
                                                     grade_list=grades))
    headers = owner()
    body = await (await client.get("/app/api/grades", headers=headers)).json()
    assert body["student"] == "Иванов Иван" and body["window_days"] == 7 and body["label_chars"] == 10
    [shown] = body["sheets"]
    assert shown["id"] == "abc123" and shown["title"] == "ДМ ПИ" and shown["label"] == "DM"
    assert shown["grades"] == [{"label": "DM HW 1", "value": 8, "raw": 8, "max": 10, "bonus": 1, "new": True}]

    woken = []
    grades.wake = lambda: woken.append(1)
    response = await client.patch("/app/api/grades", headers=headers,
                                  json={"sheets": [{"url": URL, "label": "DM"}, {"url": OTHER, "label": "Calc"}]})
    assert response.status == 200 and len((await response.json())["sheets"]) == 2 and woken == [1]
    response = await client.patch("/app/api/grades", headers=headers, json={"sheets": [{"url": "https://x.y/z"}]})
    assert response.status == 400 and "Google Sheet" in await response.text()
    assert (await client.patch("/app/api/grades", headers=headers, json={"pdf_pages": 3})).status == 400
    response = await client.patch("/app/api/grades", headers=headers, json={"window_days": 14})
    assert (await response.json())["window_days"] == 14
